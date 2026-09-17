"""
Check the bulk import: row numbers, streaming progress, and cancel.

  python verify_bulk_import.py

Three things were wrong and none of them showed up in any other suite:

  A failure reported "row 1" when the record was on spreadsheet row 3, because
  the number was the position in the parsed list, not the row in the file. On a
  1290-row upload that is the difference between finding the record and not.

  The progress bar was a timer that crept to 85% on random increments, and the
  row total came from splitting the file on newlines -- meaningless for a
  binary .xlsx.

  There was no way to stop a run.

Everything this creates is removed before it exits.
"""

import atexit
import io
import json
import os
import sys

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'enpower_skill_lab.settings')
django.setup()

from django.conf import settings                          # noqa: E402

settings.ALLOWED_HOSTS = list(settings.ALLOWED_HOSTS) + ['testserver']
settings.EMAIL_BACKEND = 'django.core.mail.backends.locmem.EmailBackend'

from django.contrib.auth import get_user_model            # noqa: E402
from django.core.files.uploadedfile import SimpleUploadedFile   # noqa: E402

from superadmin.bulk_import import SAMPLE_DATA            # noqa: E402
from verify_client import HttpsClient as Client           # noqa: E402

U = get_user_model()
PASS, FAIL = [], []
restore = {}
MARKER = 'zzbulkcheck'


def _drop(qs):
    """Delete these profiles and the login accounts behind them.

    The FK is SET_NULL, so deleting a Parent or Student on its own leaves a
    User holding the <child>-par / -stu username. The next run then resolves
    to <child>-2-par and fails on leftovers from the last one rather than on
    anything real.
    """
    ids = list(qs.filter(user__isnull=False).values_list('user_id', flat=True))
    qs.delete()
    if ids:
        U.objects.filter(id__in=ids).delete()

def _cleanup():
    from coordinator.models import ProgramCoordinator
    from parent.models import Parent
    from school_admin.models import SchoolAdmin
    from schools.models import School
    from student.models import Student
    from teacher.models import Teacher

    # Take the login accounts with the profiles. The FK is SET_NULL, so
    # deleting a Parent leaves its User behind -- and that User still holds
    # the <child>-par username, so the next run's id resolves to
    # <child>-2-par and the run fails on a leftover from the last one.
    U.objects.filter(email__startswith=MARKER).delete()
    U.objects.filter(username__startswith=MARKER).delete()
    _drop(Student.objects.filter(gr_number__startswith=MARKER))
    _drop(Parent.objects.filter(full_name__startswith=MARKER))
    Teacher.objects.filter(employee_id__startswith=MARKER).delete()
    School.objects.filter(school_code__startswith=MARKER.upper()).delete()
    for model in (SchoolAdmin, Teacher, Parent, ProgramCoordinator, Student):
        for field in ('email', 'official_email', 'school_email'):
            if any(f.name == field for f in model._meta.fields):
                model.objects.filter(**{f'{field}__startswith': MARKER}).delete()
    for pk, password in list(restore.items()):
        U.objects.filter(pk=pk).update(password=password)
    restore.clear()


atexit.register(_cleanup)


def check(label, ok, detail=''):
    (PASS if ok else FAIL).append(label)
    print(f'  {"PASS" if ok else "FAIL"}  {label}{("  - " + detail) if detail else ""}')


def sheet_with(role, rows):
    """A real downloaded sample with these data rows appended."""
    from openpyxl import load_workbook
    r = admin.get(f'/super-admin/bulk-import/{role}/sample-csv/')
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    fields = [str(cell.value or '').strip() for cell in ws[2]]
    ws.delete_rows(3, ws.max_row)
    for row in rows:
        ws.append([row.get(f, '') for f in fields])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return SimpleUploadedFile('t.xlsx', buf.read())


su = U.objects.filter(role='SUPER_ADMIN', is_active=True).first()
restore[su.pk] = U.objects.get(pk=su.pk).password
su.set_password('BulkCheck!2026')
su.save(update_fields=['password'])
admin = Client()
assert admin.login(username=su.username, password='BulkCheck!2026')

# ── the stream itself ───────────────────────────────────────────────────
print('\nSTREAMING RESPONSE')
good = {'full_name': 'ZZ Bulk A', 'email': f'{MARKER}.a@example.com',
        'phone': '9876500081', 'gender': 'Female', 'school_name': 'Nope'}
broken = {'full_name': '', 'email': f'{MARKER}.b@example.com',
          'phone': '9876500082', 'gender': 'Female', 'school_name': 'Nope'}
third = {'full_name': 'ZZ Bulk C', 'email': f'{MARKER}.c@example.com',
         'phone': '9876500083', 'gender': 'Male', 'school_name': 'Nope'}

r = admin.post('/super-admin/bulk-import/school_admin/upload-stream/',
               {'csv_file': sheet_with('school_admin', [good, broken, third])})
check('stream responds 200', r.status_code == 200, f'HTTP {r.status_code}')
check('served as ndjson', 'ndjson' in r.headers.get('Content-Type', ''),
      r.headers.get('Content-Type', ''))
check('proxies told not to buffer', r.headers.get('X-Accel-Buffering') == 'no')

events = [json.loads(line) for line in
          b''.join(r.streaming_content).decode().splitlines() if line.strip()]
rows = [e for e in events if e.get('type') == 'row']

check('opens with the row total',
      events[0].get('type') == 'start' and events[0].get('total') == 3,
      str(events[0]))
check('closes with a summary', events[-1].get('type') == 'done')
check('one event per row', len(rows) == 3, str(len(rows)))

# ── the row numbers, which is the whole complaint ───────────────────────
print('\nROW NUMBERS  (the spreadsheet\'s, not the list index)')
check('data rows are reported as 3, 4, 5',
      [e['row'] for e in rows] == [3, 4, 5], str([e['row'] for e in rows]))
failed = [e for e in rows if e['status'] == 'failed']
# The reason quotes the column heading, not the raw field name -- "Full Name",
# not "full_name" -- because the heading is what the person filling the sheet
# is looking at. Asserting the raw name would push the message back to
# something they cannot find in their spreadsheet.
check('the broken row is named by its own number, and by its column',
      any(e['row'] == 4 and 'Full Name' in (e.get('reason') or '') for e in failed),
      '; '.join(f"row {e['row']}: {e.get('reason')}" for e in failed))
# Every required column must be nameable the way the sheet names it. A
# message quoting the raw field sent a client looking for an email address
# when the column asks for a registration ID.
print(chr(10) + 'REQUIRED-COLUMN ERRORS QUOTE THE COLUMN HEADING')
from superadmin.bulk_import import EXCEL_CONFIG, _require    # noqa: E402

for _role, _cfg in EXCEL_CONFIG.items():
    labels = _cfg.get('header_map', {})
    bad = []
    for _field in sorted(_cfg.get('required_fields', ())):
        try:
            _require({}, (_field,), _role)
        except ValueError as exc:
            if str(exc).startswith(_field + ' '):
                bad.append(_field)
    check(f'{_role}: every required column is named as the sheet names it',
          not bad, ('raw field name in: ' + ', '.join(bad)) if bad else '')

try:
    _require({}, ('student_emails',), 'parent')
    _linking = ''
except ValueError as exc:
    _linking = str(exc)
check('the parent sheet says a Reg ID or GR Number will do, not just an email',
      'Reg ID' in _linking and 'GR Number' in _linking, _linking)

check('every failure carries a reason',
      all(e.get('reason') for e in failed))
check('every row carries the running counts',
      all('success' in e and 'failed' in e for e in rows))

# ── progress is countable from the stream ───────────────────────────────
print('\nPROGRESS')
check('counts only ever move forward',
      [e['success'] + e['failed'] for e in rows] == [1, 2, 3],
      str([e['success'] + e['failed'] for e in rows]))
check('the summary agrees with the rows',
      events[-1]['total'] == len(rows)
      and events[-1]['failed'] == len(failed))

# ── a rejected file still answers usefully ──────────────────────────────
print('\nREJECTED UPLOADS')
r = admin.post('/super-admin/bulk-import/school_admin/upload-stream/', {})
check('no file is refused as JSON, not a stream',
      r.status_code == 400 and 'json' in r.headers.get('Content-Type', ''),
      f'HTTP {r.status_code} {r.headers.get("Content-Type", "")}')
check('and says what was wrong',
      'No file' in json.loads(r.content).get('error', ''),
      json.loads(r.content).get('error', ''))

r = admin.post('/super-admin/bulk-import/school_admin/upload-stream/',
               {'csv_file': SimpleUploadedFile('x.txt', b'not a sheet')})
check('a wrong file type is refused', r.status_code == 400,
      f'HTTP {r.status_code}')

# ── access ──────────────────────────────────────────────────────────────
print('\nACCESS')
anon = Client()
check('anonymous cannot stream an import',
      anon.post('/super-admin/bulk-import/school_admin/upload-stream/',
                {}).status_code in (301, 302, 403, 404))

coach = U.objects.filter(role='THINKING_COACH', is_active=True).first()
if coach:
    restore[coach.pk] = U.objects.get(pk=coach.pk).password
    coach.set_password('BulkCheck!2026')
    coach.save(update_fields=['password'])
    c = Client()
    c.login(username=coach.username, password='BulkCheck!2026')
    check('a coach cannot stream a Super Admin import',
          c.post('/super-admin/bulk-import/school_admin/upload-stream/',
                 {}).status_code in (301, 302, 403, 404))

# ── the modals actually load the shared script ──────────────────────────
print('\nTHE MODALS ARE WIRED')
for url, needles in [
    ('/super-admin/students/', ['bulk-import.js', 'upload-stream', 'biCancelBtn']),
    ('/super-admin/bulk-upload/', ['bulk-import.js', 'upload-stream', 'biCancelBtn']),
]:
    r = admin.get(url, follow=True)
    if r.status_code != 200:
        check(f'{url} loads', False, f'HTTP {r.status_code}')
        continue
    body = r.content.decode(errors='ignore')
    for needle in needles:
        check(f'{url} carries {needle}', needle in body)
    check(f'{url} has no invented progress timer',
          'fakeProgress' not in body and 'readAsText' not in body)


# -- every role, not just the one this suite grew up testing ------------
# The sample workbook each role offers is the contract for that role. If it
# cannot be filled in and uploaded back, the role's import is broken however
# well the streaming machinery works.
print(chr(10) + 'EVERY ROLE ROUND-TRIPS ITS OWN SAMPLE')

ALL_ROLES = ['school', 'school_admin', 'teacher', 'student', 'parent',
             'coordinator']

from schools.models import School                          # noqa: E402

host_school = School.objects.exclude(
    school_code__startswith=MARKER.upper()).first()
check('there is a school to attach the other roles to',
      host_school is not None)

_n = [0]


def _unique(header, value, role):
    """Make the identifying columns unique so a rerun is not a duplicate.

    Everything else is left as the sample shipped it -- the point is to prove
    the sample the app hands out is one the app accepts back.
    """
    h = header.lower()
    _n[0] += 1
    n = _n[0]
    if 'email' in h:
        return f'{MARKER}.{n}@example.com'
    if h == 'school_code':
        return f'{MARKER.upper()}{n}'
    if h == 'school_name':
        return host_school.school_name if role != 'school' else f'ZZ Bulk School {n}'
    if h == 'pan_number':
        return f'ZZ{n:03d}E1234F'
    if h in ('gr_number', 'roll_number', 'employee_id', 'aadhar_number'):
        return f'{MARKER}{n}'
    if any(k in h for k in ('phone', 'mobile')) and str(value or '').strip():
        return f'90000{n:05d}'
    return value


if host_school:
    from openpyxl import load_workbook                     # noqa: E402

    for role in ALL_ROLES:
        r = admin.get(f'/super-admin/bulk-import/{role}/sample-csv/')
        if r.status_code != 200:
            check(f'{role}: sample downloads', False, f'HTTP {r.status_code}')
            continue
        check(f'{role}: sample downloads', True)

        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        fields = [str(c.value or '').strip() for c in ws[2]]
        sample = [[c.value for c in row] for row in ws.iter_rows(min_row=3)
                  if any(c.value not in (None, '') for c in row)]
        check(f'{role}: the sample ships a filled-in example row', bool(sample),
              '' if sample else 'an empty sample teaches the user nothing')
        if not sample:
            continue

        # One row per role: several of these carry uniqueness rules of their
        # own (one active admin per school, for instance) that a second copy
        # of the same sample row would trip for reasons of its own.
        row = [_unique(fields[i], v, role) if i < len(fields) else v
               for i, v in enumerate(sample[0])]
        ws.delete_rows(3, ws.max_row)
        ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        up = admin.post(f'/super-admin/bulk-import/{role}/upload-stream/',
                        {'csv_file': SimpleUploadedFile('t.xlsx', buf.read())})
        if up.status_code != 200:
            check(f'{role}: upload is accepted', False, f'HTTP {up.status_code}')
            continue

        events = [json.loads(line) for line in
                  b''.join(up.streaming_content).decode().splitlines()
                  if line.strip()]
        rows = [e for e in events if e.get('type') == 'row']
        failed = [e for e in rows if e.get('status') == 'failed']
        check(f'{role}: the row is imported', bool(rows) and not failed,
              '' if (rows and not failed) else
              ('; '.join(f"row {e['row']}: {e.get('reason')}" for e in failed)
               or 'no row event'))
        check(f'{role}: it is reported on spreadsheet row 3',
              bool(rows) and rows[0]['row'] == 3,
              str(rows[0]['row']) if rows else '-')

# -- a duplicate is refused, and says which row and why -----------------
# The complaint that started this was an unusable failure message. Uploading
# the same sheet twice is the cheapest way to produce a real one.
print(chr(10) + 'A REJECTED ROW NAMES ITSELF')
if host_school:
    r = admin.get('/super-admin/bulk-import/student/sample-csv/')
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    fields = [str(c.value or '').strip() for c in ws[2]]
    sample = [[c.value for c in row] for row in ws.iter_rows(min_row=3)
              if any(c.value not in (None, '') for c in row)][0]
    row = [_unique(fields[i], v, 'student') if i < len(fields) else v
           for i, v in enumerate(sample)]
    ws.delete_rows(3, ws.max_row)
    ws.append(row)

    def _send():
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = admin.post('/super-admin/bulk-import/student/upload-stream/',
                          {'csv_file': SimpleUploadedFile('t.xlsx', buf.read())})
        return [json.loads(l) for l in
                b''.join(resp.streaming_content).decode().splitlines()
                if l.strip()]

    first = [e for e in _send() if e.get('type') == 'row']
    second = [e for e in _send() if e.get('type') == 'row']
    check('the same sheet twice: the first goes in',
          bool(first) and first[0]['status'] != 'failed',
          '' if (first and first[0]['status'] != 'failed')
          else (str(first[0].get('reason')) if first else 'no row event'))
    dupes = [e for e in second if e.get('status') == 'failed']
    check('and the second is refused', bool(dupes),
          '' if dupes else 'a duplicate was imported twice')
    if dupes:
        check('the refusal names the spreadsheet row', dupes[0]['row'] == 3,
              str(dupes[0]['row']))
        check('and gives a reason a person can act on',
              len(dupes[0].get('reason') or '') > 10,
              repr(dupes[0].get('reason')))


# -- a parent with no email address -------------------------------------
# The platform never emails a parent and they cannot reset a password, so an
# address the school does not have should not block an import. Two of them
# must also coexist: the column is unique, and '' collides where NULL does not.
print(chr(10) + 'A PARENT CAN BE IMPORTED WITH NO EMAIL')

from datetime import date as _date                         # noqa: E402

from accounts.onboarding_ids import student_id_for         # noqa: E402
from parent.models import Parent                           # noqa: E402
from schools.models import School as _School               # noqa: E402
from student.models import Student as _Student             # noqa: E402

check('email is no longer a required column on the parent sheet',
      'email' not in EXCEL_CONFIG['parent']['required_fields'])
check('but the student link still is',
      'student_emails' in EXCEL_CONFIG['parent']['required_fields'])

_host = _School.objects.first()
_kids = []
if _host:
    for _i, (_fn, _ln) in enumerate((('Zeta', 'Noemail'), ('Theta', 'Noemail'))):
        _k = _Student.objects.create(
            first_name=_fn, last_name=_ln, gender='Male',
            date_of_birth=_date(2013, 3, _i + 1), student_class='8',
            division='C', roll_number=str(700 + _i), academic_year='2026-2027',
            school_board='CBSE', school_email=f'{MARKER}ne{_i}@example.com',
            enrollment_date=_date(2026, 6, 1), emergency_name='X',
            emergency_relationship='Parent', emergency_mobile='9000000000',
            gr_number=f'{MARKER}ne{_i}', school=_host)
        _k.skill_lab_reg_id = student_id_for(
            _host, _fn, _ln, '8', 'C', _k.date_of_birth, '2026-2027')
        _k.save()
        _kids.append(_k)

if _kids:
    from openpyxl import load_workbook as _load            # noqa: E402

    _r = admin.get('/super-admin/bulk-import/parent/sample-csv/')
    _wb = _load(io.BytesIO(_r.content))
    _ws = _wb.active
    _fields = [str(c.value or '').strip() for c in _ws[2]]
    _base = [[c.value for c in row] for row in _ws.iter_rows(min_row=3)
             if any(c.value not in (None, '') for c in row)][0]
    _ws.delete_rows(3, _ws.max_row)
    for _i, _k in enumerate(_kids):
        _row = list(_base)
        _row[_fields.index('full_name')] = f'{MARKER} NoEmail {_i}'
        _row[_fields.index('email')] = ''                  # deliberately blank
        _row[_fields.index('mobile_number')] = f'900007000{_i}'
        _row[_fields.index('student_emails')] = _k.skill_lab_reg_id
        _ws.append(_row)
    _buf = io.BytesIO()
    _wb.save(_buf)
    _buf.seek(0)

    _up = admin.post('/super-admin/bulk-import/parent/upload-stream/',
                     {'csv_file': SimpleUploadedFile('p.xlsx', _buf.read())})
    _ev = [json.loads(l) for l in
           b''.join(_up.streaming_content).decode().splitlines() if l.strip()]
    _rows = [e for e in _ev if e.get('type') == 'row']
    _bad = [e for e in _rows if e.get('status') == 'failed']
    check('both rows import with the email column blank', len(_rows) == 2 and not _bad,
          '; '.join(f"row {e['row']}: {e.get('reason')}" for e in _bad))

    _made = list(Parent.objects.filter(full_name__startswith=f'{MARKER} NoEmail'))
    check('two blank addresses coexist on a unique column', len(_made) == 2,
          '' if len(_made)==2 else f'{len(_made)} created -- NULL, not empty string, is what lets them')
    check('the address is stored as NULL, not an empty string',
          all(p.email is None for p in _made),
          str([p.email for p in _made]))

    for _p in _made:
        check(f'{_p.parent_id}: the child-derived login still works',
              Client().login(username=_p.parent_id, password=_p.parent_id))
        check(f'{_p.parent_id}: it is the child\'s id with -par',
              _p.parent_id.endswith('-par'), _p.parent_id)

    _body = admin.get('/super-admin/parents/', follow=True).content.decode(
        errors='ignore')
    # A null renders as the word "None" unless a template filter stops it.
    check('the list does not print the word "None" for a missing address',
          '>None<' not in _body)

    # With no child matched AND no address there is nothing to sign in with.
    # That must be refused, not turned into a User nobody can use.
    _wb2 = _load(io.BytesIO(admin.get(
        '/super-admin/bulk-import/parent/sample-csv/').content))
    _ws2 = _wb2.active
    _row2 = list(_base)
    _row2[_fields.index('full_name')] = f'{MARKER} Nothing'
    _row2[_fields.index('email')] = ''
    _row2[_fields.index('mobile_number')] = '9000079999'
    _row2[_fields.index('student_emails')] = 'no-such-student-anywhere'
    _ws2.delete_rows(3, _ws2.max_row)
    _ws2.append(_row2)
    _b2 = io.BytesIO()
    _wb2.save(_b2)
    _b2.seek(0)
    _up2 = admin.post('/super-admin/bulk-import/parent/upload-stream/',
                      {'csv_file': SimpleUploadedFile('p.xlsx', _b2.read())})
    _ev2 = [json.loads(l) for l in
            b''.join(_up2.streaming_content).decode().splitlines() if l.strip()]
    _r2 = [e for e in _ev2 if e.get('type') == 'row']
    _failed2 = [e for e in _r2 if e.get('status') == 'failed']
    check('no child and no address is refused, not half-created',
          bool(_failed2), '' if _failed2 else 'a parent was created with nothing to sign in with')
    if _failed2:
        check('and the refusal names the column that would fix it',
              'Reg ID' in (_failed2[0].get('reason') or ''),
              _failed2[0].get('reason'))

    _drop(_Student.objects.filter(gr_number__startswith=f'{MARKER}ne'))
    _drop(Parent.objects.filter(full_name__startswith=MARKER))


# -- the whole journey, end to end --------------------------------------
# Student -> parent linked to that student -> Download Credentials -> log in
# with exactly what the CSV printed. Each step already passed alone; the
# journey did not, because the student sheet required a parent email while the
# parent needed the student's Reg ID. Neither could be created first.
print(chr(10) + 'STUDENT, THEN PARENT, THEN THE CREDENTIALS FILE')

import csv as _csv                                         # noqa: E402

from openpyxl import load_workbook as _load2               # noqa: E402
from parent.models import Parent as _Parent                # noqa: E402
from schools.models import School as _Sch                  # noqa: E402
from student.models import Student as _Stu                 # noqa: E402

check('the student sheet does not demand a parent that cannot exist yet',
      'parent_email' not in EXCEL_CONFIG['student']['required_fields'],
      'requiring it deadlocks student-then-parent')

_school = _Sch.objects.first()


def _one_row(role, mutate):
    r = admin.get(f'/super-admin/bulk-import/{role}/sample-csv/')
    wb = _load2(io.BytesIO(r.content))
    ws = wb.active
    fields = [str(c.value or '').strip() for c in ws[2]]
    base = [[c.value for c in row] for row in ws.iter_rows(min_row=3)
            if any(c.value not in (None, '') for c in row)][0]
    row = list(base)
    mutate(fields, row)
    ws.delete_rows(3, ws.max_row)
    ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = admin.post(f'/super-admin/bulk-import/{role}/upload-stream/',
                      {'csv_file': SimpleUploadedFile('x.xlsx', buf.read())})
    events = [json.loads(l) for l in
              b''.join(resp.streaming_content).decode().splitlines() if l.strip()]
    return [e for e in events if e.get('type') == 'row']


if _school:
    def _student(f, row):
        row[f.index('first_name')] = MARKER
        row[f.index('last_name')] = 'Journey'
        row[f.index('school_name')] = _school.school_name
        row[f.index('school_email')] = f'{MARKER}journey@example.com'
        row[f.index('gr_number')] = f'{MARKER}jr'
        row[f.index('student_class')] = '6'
        row[f.index('division')] = 'A'
        row[f.index('roll_number')] = '911'
        row[f.index('parent_email')] = ''          # deliberately blank

    _r = _one_row('student', _student)
    _bad = [e for e in _r if e.get('status') == 'failed']
    check('a student imports with Parent Email left blank', _r and not _bad,
          '; '.join(f"row {e['row']}: {e.get('reason')}" for e in _bad))

    _kid = _Stu.objects.filter(gr_number=f'{MARKER}jr').first()
    check('the student got a structured registration id',
          _kid is not None and (_kid.skill_lab_reg_id or '').endswith('-stu'),
          _kid.skill_lab_reg_id if _kid else 'no student')

    if _kid:
        def _parent(f, row):
            row[f.index('full_name')] = f'{MARKER} Journey Parent'
            row[f.index('email')] = ''             # also blank
            row[f.index('mobile_number')] = '9000091111'
            row[f.index('student_emails')] = _kid.skill_lab_reg_id

        _r2 = _one_row('parent', _parent)
        _bad2 = [e for e in _r2 if e.get('status') == 'failed']
        check('a parent imports linked by that Reg ID', _r2 and not _bad2,
              '; '.join(f"row {e['row']}: {e.get('reason')}" for e in _bad2))

        _par = _Parent.objects.filter(
            full_name=f'{MARKER} Journey Parent').first()
        check('the parent id is the child id with -par',
              _par is not None
              and _par.parent_id == _kid.skill_lab_reg_id.replace('-stu', '-par'),
              (_par.parent_id if _par else 'no parent')
              + ' vs ' + _kid.skill_lab_reg_id)

        # -- the credentials file --------------------------------------
        _resp = admin.get('/super-admin/export-credentials/')
        check('the credentials file downloads', _resp.status_code == 200,
              f'HTTP {_resp.status_code}')
        _rows = list(_csv.reader(io.StringIO(_resp.content.decode())))
        _ours = [r for r in _rows if len(r) > 6 and MARKER in r[1]]
        check('it lists both the student and the parent',
              {r[0] for r in _ours} == {'Student', 'Parent'},
              str([(r[0], r[1]) for r in _ours]))

        # -- and the credentials actually work -------------------------
        for _row in _ours:
            _ok = Client().login(username=_row[5], password=_row[6])
            check(f'{_row[0]}: the file\'s own id and password sign in', _ok,
                  f'id={_row[5]} password={_row[6]}')

        _parent_row = next((r for r in _ours if r[0] == 'Parent'), None)
        if _parent_row:
            _c = Client()
            _c.login(username=_parent_row[5], password=_parent_row[6])
            _dash = _c.get('/parent/dashboard/', follow=True)
            check('the parent dashboard opens with no redirect loop',
                  _dash.status_code == 200 and len(_dash.redirect_chain) <= 1,
                  f'HTTP {_dash.status_code}, {len(_dash.redirect_chain)} redirects')
            check('and it shows the child',
                  MARKER.encode() in _dash.content)

    # A parent with no child has a random password that was never its login
    # id. Printing the id there hands out a credential that cannot work.
    def _orphan(f, row):
        row[f.index('full_name')] = f'{MARKER} Orphan Parent'
        row[f.index('email')] = f'{MARKER}.orphan@example.com'
        row[f.index('mobile_number')] = '9000092222'
        row[f.index('student_emails')] = f'{MARKER}.orphan@example.com'

    _one_row('parent', _orphan)
    _op = _Parent.objects.filter(full_name=f'{MARKER} Orphan Parent').first()
    if _op:
        _rows = list(_csv.reader(io.StringIO(
            admin.get('/super-admin/export-credentials/').content.decode())))
        _orow = next((r for r in _rows if len(r) > 6 and 'Orphan' in r[1]), None)
        check('a parent with no child is not given a password that cannot work',
              _orow is not None and _orow[6] != _orow[5],
              str(_orow))

    _drop(_Stu.objects.filter(gr_number__startswith=MARKER))
    _drop(_Parent.objects.filter(full_name__startswith=MARKER))


# -- the school name, and what a miss says ------------------------------
# A real upload failed on all 336 rows with "School \'Saint Capitanio\' not
# found" because the registered name was longer. The lookup is right to stay
# exact -- guessing the nearest school would attach 336 students to the wrong
# one -- but the refusal has to say what the sheet should have contained.
print(chr(10) + 'A MISSING SCHOOL NAMES THE ONES THAT NEARLY MATCH')

from superadmin.bulk_import import _school_by_name          # noqa: E402

_reg = _School.objects.create(
    school_name=f'{MARKER} Capitanio High School',
    school_code=f'{MARKER.upper()}CAP', board='CBSE', school_type='Private',
    medium='English', school_email=f'{MARKER}cap@example.com',
    school_phone='9000000001', principal_name='P', principal_phone='9000000002',
    principal_email=f'{MARKER}capp@example.com', branch_address='A',
    city='Mumbai', state='MH', pincode='400001',
    emergency_contact_person='X', emergency_phone='9000000003')

check('an exact name still matches',
      _school_by_name(_reg.school_name).id == _reg.id)
check('and case does not matter',
      _school_by_name(_reg.school_name.upper()).id == _reg.id)


def _miss(name):
    try:
        _school_by_name(name)
        return ''
    except ValueError as exc:
        return str(exc)


# The shape the real upload took: the sheet holds a prefix of the registered
# name. difflib alone scores that pair poorly, which is why substring is tried
# first.
_short = _miss(f'{MARKER} Capitanio')
check('a name that is a prefix of the real one is refused, not guessed at',
      'not found' in _short, _short)
check('and the refusal names the school that was meant',
      _reg.school_name in _short, _short)

_typo = _miss(f'{MARKER} Capitano')
check('a typo gets the same help', _reg.school_name in _typo, _typo)

_none = _miss('Nothing Like Any Registered School At All')
check('an unrelated name says how to find the right one',
      'not found' in _none and 'school list' in _none, _none)
check('and does not invent a suggestion',
      'Did you mean' not in _none, _none)

# End to end: the message has to survive the import and reach the modal.
if _school:
    def _wrong_school(f, row):
        row[f.index('first_name')] = MARKER
        row[f.index('last_name')] = 'WrongSchool'
        row[f.index('school_name')] = f'{MARKER} Capitanio'   # the prefix
        row[f.index('school_email')] = f'{MARKER}ws@example.com'
        row[f.index('gr_number')] = f'{MARKER}ws'
        row[f.index('student_class')] = '6'
        row[f.index('division')] = 'A'
        row[f.index('roll_number')] = '912'
        row[f.index('parent_email')] = ''

    _wr = _one_row('student', _wrong_school)
    _wf = [e for e in _wr if e.get('status') == 'failed']
    check('the row is refused rather than filed under the wrong school',
          bool(_wf), '' if _wf else 'a student was created against a name that does not exist')
    if _wf:
        check('and the suggestion reaches the failure list the user reads',
              _reg.school_name in (_wf[0].get('reason') or ''),
              _wf[0].get('reason'))

_reg.delete()

_cleanup()

print('\n' + '=' * 62)
print(f'PASS {len(PASS)}   FAIL {len(FAIL)}')
for f in FAIL:
    print('  FAILED:', f)
sys.exit(1 if FAIL else 0)
